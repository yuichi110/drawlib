# Copyright (c) 2026 Yuichi Ito (yuichi@yuichi.com)
#
# This software is licensed under the Apache License, Version 2.0.
# For more information, please visit: https://github.com/yuichi110/drawlib
#
# This software is provided "as is", without warranty of any kind,
# express or implied, including but not limited to the warranties of
# merchantability, fitness for a particular purpose and noninfringement.

"""MS Word (.docx) document exporter from Markdown AST and drawlib diagrams."""

from __future__ import annotations

import os
import re
import sys
import tempfile
from collections.abc import Callable
from typing import TYPE_CHECKING, Any

from markdown_it import MarkdownIt
from markdown_it.token import Token
from pygments import lex

from drawlib._builder._common.cache import BuildImageCache
from drawlib._builder._common.office import (
    ensure_static_office_image,
    normalize_opc_zip_timestamps,
    resolve_office_datetime,
)
from drawlib._builder.doc_builder.compiler.base import validate_markdown_links
from drawlib._builder.doc_builder.detector import detect_document_type
from drawlib._builder.doc_builder.highlight import resolve_lexer
from drawlib._builder.doc_builder.merger import _extract_title, expand_input_files
from drawlib._builder.doc_builder.processor import DrawlibBlockProcessor
from drawlib._builder.doc_builder.processor.options import (
    parse_block_info,
    resolve_block_image_paths,
)
from drawlib._builder.doc_builder.progress import (
    FileBuildProgress,
    check_document_output_duplicates,
    format_duplicate_output_error,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.text.paragraph import Paragraph

    from drawlib._builder.doc_builder._docx_xml import DocxTheme

_USABLE_WIDTH_MM: float = 178.0
_MAX_IMAGE_HEIGHT_MM: float = 220.0
_HTML_PROSE_REF_PX: float = 760.0


def _sanitize_bookmark_name(raw: str, fallback: str) -> str:
    """Normalize a string into a valid Word OpenXML bookmark name (<= 40 chars, [A-Za-z0-9_])."""
    cleaned = re.sub(r"[^a-zA-Z0-9_]+", "_", raw).strip("_")
    if not cleaned or not cleaned[0].isalpha():
        cleaned = f"bm_{cleaned}" if cleaned else fallback
    return cleaned[:40]


def _is_drawlib_fence(info: str) -> bool:
    """Return True if a fence info string represents a ```drawlib block."""
    stripped = info.strip()
    if not stripped:
        return False
    first_word = stripped.split()[0]
    return first_word == "drawlib"


def _compute_image_width_mm(
    px_w: int,
    px_h: int,
    width_spec: str | None,
    *,
    is_diagram: bool,
) -> float:
    """Calculate target image width in millimeters fitted to A4 body margins."""
    target_mm = _USABLE_WIDTH_MM
    if width_spec:
        spec = width_spec.strip().lower()
        try:
            if spec.endswith("%"):
                pct = float(spec[:-1].strip())
                target_mm = _USABLE_WIDTH_MM * max(0.05, min(pct / 100.0, 1.0))
            elif spec.endswith("px"):
                px_val = float(spec[:-2].strip())
                target_mm = _USABLE_WIDTH_MM * (px_val / _HTML_PROSE_REF_PX)
            elif spec.endswith("mm"):
                target_mm = float(spec[:-2].strip())
            elif spec.endswith("cm"):
                target_mm = float(spec[:-2].strip()) * 10.0
            elif spec.endswith("in"):
                target_mm = float(spec[:-2].strip()) * 25.4
            else:
                px_val = float(spec)
                target_mm = _USABLE_WIDTH_MM * (px_val / _HTML_PROSE_REF_PX)
        except ValueError:
            target_mm = _USABLE_WIDTH_MM
    elif not is_diagram and px_w > 0:
        # Keep small static icons/badges proportional instead of stretching to 178mm
        natural_mm = px_w * 25.4 / 144.0
        target_mm = min(_USABLE_WIDTH_MM, natural_mm)

    target_mm = max(12.0, min(_USABLE_WIDTH_MM, target_mm))
    if px_w > 0 and px_h > 0:
        est_height_mm = target_mm * (px_h / px_w)
        if est_height_mm > _MAX_IMAGE_HEIGHT_MM:
            target_mm = _MAX_IMAGE_HEIGHT_MM * (px_w / px_h)
    return target_mm


def _resolve_link_target(
    href: str,
    chapter_anchors: dict[str, str],
    heading_anchors: dict[str, str],
) -> tuple[str | None, str | None]:
    """Resolve an href into (external_url, internal_bookmark_anchor)."""
    clean = href.strip()
    if not clean:
        return (None, None)
    if clean.startswith(("http://", "https://", "mailto:")):
        return (clean, None)
    if clean.startswith("#"):
        frag = clean[1:].strip()
        anchor = heading_anchors.get(frag) or chapter_anchors.get(frag) or _sanitize_bookmark_name(frag, "bm_ref")
        return (None, anchor)

    file_part = os.path.basename(clean.split("#", 1)[0])
    return (None, chapter_anchors.get(file_part))


def _parse_cell_align(token: Token) -> str | None:
    """Extract 'left', 'center', or 'right' alignment from a th_open/td_open token."""
    style_attr = token.attrGet("style")
    if style_attr:
        m = re.search(r"text-align\s*:\s*(left|center|right)", str(style_attr), re.IGNORECASE)
        if m:
            return m.group(1).lower()
    align_attr = token.attrGet("align")
    if align_attr and str(align_attr).lower() in {"left", "center", "right"}:
        return str(align_attr).lower()
    return None


class _DocxChapterRenderer:
    """Stateful converter from MarkdownIt tokens into a python-docx Document."""

    def __init__(
        self,
        doc: DocxDocument,
        theme: DocxTheme,
        scratch_dir: str,
        chapter_anchors: dict[str, str],
    ) -> None:
        self.doc = doc
        self.theme = theme
        self.scratch_dir = scratch_dir
        self.chapter_anchors = chapter_anchors
        self.heading_anchors: dict[str, str] = {}
        self._bookmark_id_seq = 1
        self._md = MarkdownIt("gfm-like", {"html": True, "linkify": False, "typographer": True})

    def next_bookmark_id(self) -> int:
        """Allocate a unique OpenXML bookmark integer ID."""
        bid = self._bookmark_id_seq
        self._bookmark_id_seq += 1
        return bid

    def _attach_bookmark(self, paragraph: Paragraph, bookmark_name: str) -> None:
        from drawlib._builder.doc_builder._docx_xml import add_bookmark

        add_bookmark(paragraph, bookmark_name, self.next_bookmark_id())

    def _emit_image_block(
        self,
        img_path: str,
        *,
        width_spec: str | None = None,
        align: str | None = "center",
        caption: str | None = None,
        is_diagram: bool = False,
        pending_chapter_bookmark: str | None = None,
    ) -> bool:
        """Embed a static raster image paragraph and optional caption paragraph."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm, Pt

        if not os.path.isfile(img_path):
            return False

        static_path, px_w, px_h = ensure_static_office_image(img_path, self.scratch_dir)
        width_mm = _compute_image_width_mm(px_w, px_h, width_spec, is_diagram=is_diagram)

        p = self.doc.add_paragraph()
        if pending_chapter_bookmark:
            self._attach_bookmark(p, pending_chapter_bookmark)

        align_norm = (align or "center").lower()
        if align_norm == "left":
            target_align = WD_ALIGN_PARAGRAPH.LEFT
        elif align_norm == "right":
            target_align = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            target_align = WD_ALIGN_PARAGRAPH.CENTER
        p.alignment = target_align

        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2 if caption else 8)
        if caption:
            p.paragraph_format.keep_with_next = True

        run = p.add_run()
        run.add_picture(static_path, width=Mm(width_mm))

        if caption:
            cap_style = "Caption" if "Caption" in self.doc.styles else "Normal"
            cap_p = self.doc.add_paragraph(style=cap_style)
            cap_p.alignment = target_align
            cap_run = cap_p.add_run(caption)
            from drawlib._builder.doc_builder._docx_xml import set_run_fonts

            set_run_fonts(cap_run, self.theme.latin_font, self.theme.east_asia_font)
        return True

    def _emit_code_block(
        self,
        code: str,
        lang: str,
        *,
        pending_chapter_bookmark: str | None = None,
    ) -> bool:
        """Render a syntax-highlighted fenced code block inside a 1x1 shaded Word table."""
        from docx.shared import Pt, RGBColor

        from drawlib._builder.doc_builder._docx_xml import (
            add_post_table_spacer,
            resolve_pygments_style,
            set_run_fonts,
            set_table_full_width,
            set_table_row_flags,
            style_table_cell,
        )

        table = self.doc.add_table(rows=1, cols=1)
        set_table_full_width(table, _USABLE_WIDTH_MM)
        set_table_row_flags(table.rows[0], is_header=False, cant_split=False)
        cell = table.cell(0, 0)
        style_table_cell(
            cell,
            fill_hex=self.theme.code_bg_hex,
            border_hex=self.theme.border_hex,
            all_borders=True,
            is_header=False,
            pad_v_dxa=130,
            pad_h_dxa=160,
        )

        p = cell.paragraphs[0]
        if pending_chapter_bookmark:
            self._attach_bookmark(p, pending_chapter_bookmark)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.25

        clean_code = code.rstrip("\r\n")
        lexer = resolve_lexer(lang, clean_code)
        tokens_list = list(lex(clean_code, lexer))
        if tokens_list and tokens_list[-1][1].endswith("\n"):
            last_ttype, last_val = tokens_list[-1]
            trimmed = last_val[:-1]
            if trimmed:
                tokens_list[-1] = (last_ttype, trimmed)
            else:
                tokens_list.pop()

        for ttype, val in tokens_list:
            if not val:
                continue
            rgb, is_bold, is_italic = resolve_pygments_style(ttype, self.theme)
            parts = val.split("\n")
            for idx, part in enumerate(parts):
                if part:
                    run = p.add_run(part)
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(*rgb)
                    if is_bold:
                        run.bold = True
                    if is_italic:
                        run.italic = True
                    set_run_fonts(run, self.theme.mono_font, self.theme.east_asia_font)
                if idx < len(parts) - 1:
                    p.add_run().add_break()

        add_post_table_spacer(self.doc)
        return bool(pending_chapter_bookmark)

    def _render_inline(
        self,
        paragraph: Paragraph,
        children: list[Token] | None,
        src_dir: str,
        *,
        base_bold: bool = False,
        base_color: tuple[int, int, int] | None = None,
        font_size_pt: float | None = None,
    ) -> None:
        """Render markdown-it inline child tokens into a Word paragraph."""
        from docx.shared import Mm, Pt, RGBColor

        from drawlib._builder.doc_builder._docx_xml import (
            add_hyperlink_run,
            set_run_fonts,
            set_run_shading,
        )

        if not children:
            return

        bold_depth = 1 if base_bold else 0
        italic_depth = 0
        strike_depth = 0
        active_href: str | None = None
        eff_color = base_color if base_color is not None else self.theme.text_color

        for child in children:
            ctype = child.type
            if ctype == "strong_open":
                bold_depth += 1
            elif ctype == "strong_close":
                bold_depth = max(1 if base_bold else 0, bold_depth - 1)
            elif ctype == "em_open":
                italic_depth += 1
            elif ctype == "em_close":
                italic_depth = max(0, italic_depth - 1)
            elif ctype == "s_open":
                strike_depth += 1
            elif ctype == "s_close":
                strike_depth = max(0, strike_depth - 1)
            elif ctype == "link_open":
                href_val = child.attrGet("href")
                active_href = str(href_val) if href_val else None
            elif ctype == "link_close":
                active_href = None
            elif ctype == "softbreak":
                paragraph.add_run(" ")
            elif ctype == "hardbreak":
                paragraph.add_run().add_break()
            elif ctype == "html_inline":
                raw_html = child.content.strip().lower()
                if raw_html.startswith("<br"):
                    paragraph.add_run().add_break()
            elif ctype == "code_inline":
                if active_href:
                    ext_url, int_anchor = _resolve_link_target(
                        active_href, self.chapter_anchors, self.heading_anchors
                    )
                    if ext_url or int_anchor:
                        add_hyperlink_run(
                            paragraph,
                            child.content,
                            theme=self.theme,
                            url=ext_url,
                            anchor=int_anchor,
                            bold=bold_depth > 0,
                            italic=italic_depth > 0,
                            strikethrough=strike_depth > 0,
                            is_code=True,
                        )
                        continue
                run = paragraph.add_run(child.content)
                set_run_fonts(run, self.theme.mono_font, self.theme.east_asia_font)
                run.font.size = Pt((font_size_pt * 0.9) if font_size_pt else 9.5)
                run.font.color.rgb = RGBColor(*(base_color or self.theme.inline_code_color))
                if bold_depth > 0:
                    run.bold = True
                if italic_depth > 0:
                    run.italic = True
                if strike_depth > 0:
                    run.font.strike = True
                set_run_shading(run, self.theme.inline_code_bg_hex)
            elif ctype == "image":
                src_attr = child.attrGet("src")
                if src_attr:
                    clean_ref = str(src_attr).split("#")[0].split("?")[0]
                    abs_img = os.path.abspath(os.path.join(src_dir, clean_ref))
                    if os.path.isfile(abs_img):
                        static_path, px_w, px_h = ensure_static_office_image(abs_img, self.scratch_dir)
                        w_mm = _compute_image_width_mm(px_w, px_h, None, is_diagram=False)
                        paragraph.add_run().add_picture(static_path, width=Mm(w_mm))
            elif ctype == "text":
                if not child.content:
                    continue
                if active_href:
                    ext_url, int_anchor = _resolve_link_target(
                        active_href, self.chapter_anchors, self.heading_anchors
                    )
                    if ext_url or int_anchor:
                        add_hyperlink_run(
                            paragraph,
                            child.content,
                            theme=self.theme,
                            url=ext_url,
                            anchor=int_anchor,
                            bold=bold_depth > 0,
                            italic=italic_depth > 0,
                            strikethrough=strike_depth > 0,
                            is_code=False,
                        )
                        continue
                run = paragraph.add_run(child.content)
                set_run_fonts(run, self.theme.latin_font, self.theme.east_asia_font)
                if font_size_pt is not None:
                    run.font.size = Pt(font_size_pt)
                run.font.color.rgb = RGBColor(*eff_color)
                if bold_depth > 0:
                    run.bold = True
                if italic_depth > 0:
                    run.italic = True
                if strike_depth > 0:
                    run.font.strike = True

    def _emit_table(
        self,
        tokens: list[Token],
        start_idx: int,
        src_dir: str,
        *,
        pending_chapter_bookmark: str | None = None,
    ) -> tuple[int, bool]:
        """Consume table_open..table_close tokens and emit a styled native Word table."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        from drawlib._builder.doc_builder._docx_xml import (
            add_post_table_spacer,
            set_table_full_width,
            set_table_row_flags,
            style_table_cell,
        )

        rows_data: list[tuple[bool, list[tuple[str | None, Token | None]]]] = []
        in_thead = False
        current_cells: list[tuple[str | None, Token | None]] = []
        current_align: str | None = None
        current_inline: Token | None = None

        i = start_idx + 1
        while i < len(tokens):
            tok = tokens[i]
            tt = tok.type
            if tt == "table_close":
                i += 1
                break
            if tt == "thead_open":
                in_thead = True
            elif tt == "thead_close":
                in_thead = False
            elif tt == "tr_open":
                current_cells = []
            elif tt == "tr_close":
                if current_cells:
                    rows_data.append((in_thead, current_cells))
            elif tt in {"th_open", "td_open"}:
                current_align = _parse_cell_align(tok)
                current_inline = None
            elif tt == "inline":
                current_inline = tok
            elif tt in {"th_close", "td_close"}:
                current_cells.append((current_align, current_inline))
            i += 1

        if not rows_data:
            return (i, False)

        num_cols = max((len(cells) for _, cells in rows_data), default=1)
        table = self.doc.add_table(rows=len(rows_data), cols=num_cols)
        set_table_full_width(table, _USABLE_WIDTH_MM)
        table.autofit = True

        bookmark_attached = False
        for r_idx, (is_header, cells) in enumerate(rows_data):
            row_obj = table.rows[r_idx]
            set_table_row_flags(row_obj, is_header=is_header, cant_split=True)
            for c_idx in range(num_cols):
                cell = table.cell(r_idx, c_idx)
                style_table_cell(
                    cell,
                    fill_hex=self.theme.surface_bg_hex if is_header else None,
                    border_hex=self.theme.border_hex,
                    all_borders=False,
                    is_header=is_header,
                )
                p = cell.paragraphs[0]
                if pending_chapter_bookmark and not bookmark_attached:
                    self._attach_bookmark(p, pending_chapter_bookmark)
                    bookmark_attached = True
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.25

                if c_idx < len(cells):
                    align_str, inline_tok = cells[c_idx]
                    if align_str == "center":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif align_str == "right":
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    if inline_tok is not None:
                        self._render_inline(
                            p,
                            inline_tok.children,
                            src_dir,
                            base_bold=is_header,
                            base_color=self.theme.heading_color if is_header else self.theme.text_color,
                            font_size_pt=9.5,
                        )

        add_post_table_spacer(self.doc)
        return (i, bookmark_attached)

    def _handle_html_block(
        self,
        html_content: str,
        src_dir: str,
        *,
        pending_chapter_bookmark: str | None = None,
    ) -> bool:
        """Handle common raw HTML blocks in Markdown (page-break divs and <img> tags)."""
        lowered = html_content.lower()
        if "page-break" in lowered or "break-after: page" in lowered or "break-before: page" in lowered:
            self.doc.add_page_break()
            return False

        img_match = re.search(r'<img\b[^>]*?\bsrc=["\']([^"\']+)["\']([^>]*)>', html_content, re.IGNORECASE)
        if img_match:
            raw_src = img_match.group(1).strip()
            attrs_tail = html_content
            width_match = re.search(r'\bwidth=["\']?([0-9.]+(?:%|px|mm|cm|in)?)["\']?', attrs_tail, re.IGNORECASE)
            style_w_match = re.search(r"width\s*:\s*([0-9.]+(?:%|px|mm|cm|in))", attrs_tail, re.IGNORECASE)
            width_spec = (
                style_w_match.group(1)
                if style_w_match
                else (width_match.group(1) if width_match else None)
            )
            cap_match = re.search(r"<figcaption[^>]*>(.*?)</figcaption>", html_content, re.DOTALL | re.IGNORECASE)
            caption = re.sub(r"<[^>]+>", "", cap_match.group(1)).strip() if cap_match else None
            clean_ref = raw_src.split("#")[0].split("?")[0]
            abs_img = os.path.abspath(os.path.join(src_dir, clean_ref))
            return self._emit_image_block(
                abs_img,
                width_spec=width_spec,
                align="center",
                caption=caption,
                is_diagram=False,
                pending_chapter_bookmark=pending_chapter_bookmark,
            )
        return False

    def render_chapter(
        self,
        markdown_text: str,
        *,
        src_abs: str,
        chapter_bookmark: str,
        processor: DrawlibBlockProcessor | None,
        dpi: int,
        progress_callback: Callable[[int, int, bool], None] | None = None,
    ) -> None:
        """Parse and append a single Markdown chapter into the active Word Document."""
        from docx.shared import Mm

        from drawlib._builder.doc_builder._docx_xml import (
            allocate_ordered_list_num_id,
            apply_ordered_numbering,
            style_blockquote_paragraph,
            style_hr_paragraph,
        )

        src_dir = os.path.dirname(src_abs)
        doc_base_name = os.path.splitext(os.path.basename(src_abs))[0]
        tokens = self._md.parse(markdown_text)

        total_drawlib = sum(
            1 for t in tokens if t.type == "fence" and _is_drawlib_fence(t.info)
        )
        drawlib_counter = 0

        # Pre-register heading slug anchors within this chapter
        for idx_t, tok in enumerate(tokens):
            if tok.type == "heading_open" and idx_t + 1 < len(tokens) and tokens[idx_t + 1].type == "inline":
                raw_h = tokens[idx_t + 1].content.strip()
                slug = re.sub(r"[^a-z0-9_-]+", "-", raw_h.lower()).strip("-")
                if slug:
                    self.heading_anchors[slug] = _sanitize_bookmark_name(f"{chapter_bookmark}_{slug}", chapter_bookmark)

        bookmark_pending: str | None = chapter_bookmark
        list_stack: list[tuple[str, int | None]] = []
        list_item_para_counts: list[int] = []
        blockquote_depth = 0

        i = 0
        while i < len(tokens):
            tok = tokens[i]
            tt = tok.type

            if tt == "heading_open":
                lvl = max(1, min(6, int(tok.tag[1])))
                inline_tok = tokens[i + 1] if (i + 1 < len(tokens) and tokens[i + 1].type == "inline") else None
                para = self.doc.add_paragraph(style=f"Heading {lvl}")
                if bookmark_pending:
                    self._attach_bookmark(para, bookmark_pending)
                    bookmark_pending = None
                if inline_tok is not None:
                    slug = re.sub(r"[^a-z0-9_-]+", "-", inline_tok.content.strip().lower()).strip("-")
                    if slug and slug in self.heading_anchors:
                        self._attach_bookmark(para, self.heading_anchors[slug])
                    h_color = self.theme.muted_color if lvl == 6 else self.theme.heading_color
                    self._render_inline(
                        para,
                        inline_tok.children,
                        src_dir,
                        base_bold=True,
                        base_color=h_color,
                    )
                i += 3
                continue

            if tt == "bullet_list_open":
                list_stack.append(("bullet", None))
                i += 1
                continue

            if tt == "bullet_list_close":
                if list_stack:
                    list_stack.pop()
                i += 1
                continue

            if tt == "ordered_list_open":
                start_attr = tok.attrGet("start")
                start_num = int(start_attr) if (start_attr and str(start_attr).isdigit()) else 1
                num_id = allocate_ordered_list_num_id(self.doc, start=start_num)
                list_stack.append(("ordered", num_id))
                i += 1
                continue

            if tt == "ordered_list_close":
                if list_stack:
                    list_stack.pop()
                i += 1
                continue

            if tt == "list_item_open":
                list_item_para_counts.append(0)
                i += 1
                continue

            if tt == "list_item_close":
                if list_item_para_counts:
                    list_item_para_counts.pop()
                i += 1
                continue

            if tt == "blockquote_open":
                blockquote_depth += 1
                i += 1
                continue

            if tt == "blockquote_close":
                blockquote_depth = max(0, blockquote_depth - 1)
                i += 1
                continue

            if tt == "hr":
                hr_p = self.doc.add_paragraph()
                style_hr_paragraph(hr_p, self.theme.border_hex)
                i += 1
                continue

            if tt == "paragraph_open":
                inline_tok = tokens[i + 1] if (i + 1 < len(tokens) and tokens[i + 1].type == "inline") else None

                # Standalone single Markdown image paragraph: ![caption](path.png)
                if (
                    not list_stack
                    and blockquote_depth == 0
                    and inline_tok is not None
                    and inline_tok.children is not None
                    and len(inline_tok.children) == 1
                    and inline_tok.children[0].type == "image"
                ):
                    img_tok = inline_tok.children[0]
                    src_attr = img_tok.attrGet("src")
                    alt_text = img_tok.content.strip() or None
                    if src_attr:
                        clean_ref = str(src_attr).split("#")[0].split("?")[0]
                        abs_img = os.path.abspath(os.path.join(src_dir, clean_ref))
                        attached = self._emit_image_block(
                            abs_img,
                            align="center",
                            caption=alt_text,
                            is_diagram=False,
                            pending_chapter_bookmark=bookmark_pending,
                        )
                        if attached:
                            bookmark_pending = None
                        i += 3
                        continue

                if list_stack:
                    list_kind, list_num_id = list_stack[-1]
                    depth_idx = len(list_stack) - 1
                    item_p_idx = list_item_para_counts[-1] if list_item_para_counts else 0
                    if list_item_para_counts:
                        list_item_para_counts[-1] = item_p_idx + 1

                    if item_p_idx == 0:
                        if list_kind == "bullet":
                            b_styles = ("List Bullet", "List Bullet 2", "List Bullet 3")
                            b_style = b_styles[min(depth_idx, len(b_styles) - 1)]
                            para = self.doc.add_paragraph(
                                style=b_style if b_style in self.doc.styles else "Normal"
                            )
                        else:
                            para = self.doc.add_paragraph(
                                style="List Number" if "List Number" in self.doc.styles else "Normal"
                            )
                            if list_num_id is not None:
                                apply_ordered_numbering(para, list_num_id, depth_idx)
                    else:
                        para = self.doc.add_paragraph(style="Normal")
                        para.paragraph_format.left_indent = Mm(6.5 * (depth_idx + 1))
                elif blockquote_depth > 0:
                    para = self.doc.add_paragraph(style="Normal")
                    style_blockquote_paragraph(para, self.theme, depth=blockquote_depth)
                else:
                    para = self.doc.add_paragraph(style="Normal")

                if bookmark_pending:
                    self._attach_bookmark(para, bookmark_pending)
                    bookmark_pending = None

                if inline_tok is not None:
                    self._render_inline(para, inline_tok.children, src_dir)

                i += 3
                continue

            if tt in {"fence", "code_block"}:
                if tt == "fence" and _is_drawlib_fence(tok.info):
                    drawlib_counter += 1
                    info_rest = tok.info.strip()[len("drawlib") :].strip()
                    options = parse_block_info(info_rest)
                    line_no = (tok.map[0] + 1) if tok.map else 1

                    _, canonical_target = resolve_block_image_paths(
                        options=options,
                        doc_base_name=doc_base_name,
                        block_counter=drawlib_counter,
                        default_format="png",
                        output_dir=None,
                        require_file=True,
                        line_number=line_no,
                    )
                    cand_ext = os.path.splitext(canonical_target)[1].lower()
                    render_ext = cand_ext if cand_ext in {".png", ".webp", ".apng"} else ".png"
                    tmp_render_path = os.path.join(
                        self.scratch_dir,
                        f"{doc_base_name}_{drawlib_counter}_{os.getpid()}{render_ext}",
                    )

                    if processor is not None:
                        processor.render_block_to_file(
                            tok.content.strip(),
                            tmp_render_path,
                            source_filename=src_abs,
                            cache_target_file=canonical_target,
                            no_cache=options.no_cache,
                            dpi=dpi,
                        )
                    if progress_callback is not None:
                        progress_callback(drawlib_counter, total_drawlib, drawlib_counter == total_drawlib)

                    if options.code == "show":
                        attached = self._emit_code_block(
                            tok.content,
                            "python",
                            pending_chapter_bookmark=bookmark_pending,
                        )
                        if attached:
                            bookmark_pending = None

                    attached_img = self._emit_image_block(
                        tmp_render_path,
                        width_spec=options.width,
                        align=options.align or "center",
                        caption=options.caption,
                        is_diagram=True,
                        pending_chapter_bookmark=bookmark_pending,
                    )
                    if attached_img:
                        bookmark_pending = None
                else:
                    lang = tok.info.strip().split()[0] if tok.info.strip() else ""
                    attached = self._emit_code_block(
                        tok.content,
                        lang,
                        pending_chapter_bookmark=bookmark_pending,
                    )
                    if attached:
                        bookmark_pending = None
                i += 1
                continue

            if tt == "table_open":
                i, attached_tbl = self._emit_table(
                    tokens,
                    i,
                    src_dir,
                    pending_chapter_bookmark=bookmark_pending,
                )
                if attached_tbl:
                    bookmark_pending = None
                continue

            if tt == "html_block":
                attached_html = self._handle_html_block(
                    tok.content,
                    src_dir,
                    pending_chapter_bookmark=bookmark_pending,
                )
                if attached_html:
                    bookmark_pending = None
                i += 1
                continue

            i += 1

        if bookmark_pending:
            fallback_p = self.doc.add_paragraph()
            self._attach_bookmark(fallback_p, bookmark_pending)


def _emit_toc_section(
    doc: DocxDocument,
    theme: DocxTheme,
    toc_entries: list[tuple[str, str]],
    *,
    page_break: bool,
) -> None:
    """Insert a clickable Table of Contents section linking to chapter bookmarks."""
    from drawlib._builder.doc_builder._docx_xml import add_hyperlink_run

    if not toc_entries:
        return

    if page_break:
        doc.add_page_break()

    toc_h = doc.add_paragraph(style="Heading 2")
    run_h = toc_h.add_run("Table of Contents")
    from drawlib._builder.doc_builder._docx_xml import set_run_fonts

    set_run_fonts(run_h, theme.heading_font, theme.east_asia_font)

    b_style = "List Bullet" if "List Bullet" in doc.styles else "Normal"
    for bm_name, chap_title in toc_entries:
        item_p = doc.add_paragraph(style=b_style)
        add_hyperlink_run(
            item_p,
            chap_title,
            theme=theme,
            anchor=bm_name,
        )


def export_chapters_to_docx(
    inputs: list[str],
    output_path: str,
    *,
    title: str | None = None,
    page_break: bool = True,
    generate_index: bool = False,
    dpi: int = 200,
    styles_path: str | None = None,
    utils_path: str | None = None,
    css_path: str | None = None,
    no_cache: bool = False,
    timestamp: bool = False,
) -> list[str]:
    """Compile and merge Markdown inputs into a single native MS Word (.docx) document.

    Args:
        inputs: Ordered list of input Markdown file or directory paths.
        output_path: Destination .docx file path.
        title: Optional document title override for Word core_properties.
        page_break: Insert page breaks between merged chapters. Defaults to True.
        generate_index: Insert a Table of Contents after the 1st chapter. Defaults to False.
        dpi: Rasterization DPI for embedded drawlib diagrams. Defaults to 200.
        styles_path: Optional Python styles.py script path.
        utils_path: Optional Python utils.py script path.
        css_path: Optional project style.css path for theme extraction.
        no_cache: Disable reading/writing the SQLite build image cache.
        timestamp: Preserve current build timestamp instead of normalizing it.

    Returns:
        list[str]: Ordered list of compiled absolute Markdown file paths.
    """
    from drawlib._builder.doc_builder._docx_xml import (
        configure_docx_document,
        import_docx,
        parse_docx_theme,
    )

    if not inputs:
        raise ValueError("At least one input file or directory must be specified.")
    if dpi <= 0:
        raise ValueError(f"dpi must be a positive integer, got {dpi}.")

    docx = import_docx()

    raw_files = expand_input_files(inputs)
    file_list = [f for f in raw_files if os.path.splitext(f)[1].lower() in {".md", ".markdown"}]
    if not file_list:
        raise ValueError("No Markdown (.md) files found in the specified inputs for Word (.docx) export.")

    if len(inputs) == 1 and os.path.isdir(os.path.abspath(inputs[0])):
        base_root = os.path.abspath(inputs[0])
    elif len(file_list) > 1:
        common = os.path.commonpath(file_list)
        base_root = common if os.path.isdir(common) else os.path.dirname(common)
    else:
        base_root = os.path.dirname(file_list[0])

    if css_path is None:
        c_cand = os.path.join(base_root, "style.css")
        if os.path.isfile(c_cand):
            css_path = c_cand

    display_names = ["/" + os.path.relpath(s_abs, base_root).replace(os.sep, "/") for s_abs in file_list]
    seen_sources: dict[str, str] = {}
    for s_abs, disp_name in zip(file_list, display_names):
        if s_abs in seen_sources:
            raise ValueError(format_duplicate_output_error(s_abs, seen_sources[s_abs], disp_name))
        seen_sources[s_abs] = disp_name

    max_blocks = check_document_output_duplicates(
        tasks=[(s_abs, "", True) for s_abs in file_list],
        display_names=display_names,
        image_format="png",
        embed_images=False,
    )
    name_width = max((len(n) for n in display_names), default=0)
    image_width = len(str(max(max_blocks, 0)))

    # Pre-read chapters so titles, links, and TOC entries are known upfront
    chapters_meta: list[tuple[str, str, str, str, Any]] = []
    chapter_anchors: dict[str, str] = {}
    inferred_title: str | None = title

    for idx, (src_abs, disp_name) in enumerate(zip(file_list, display_names), start=1):
        with open(src_abs, "r", encoding="utf-8") as f:
            content = f.read()
        validate_markdown_links(src_abs, content)
        doc_info = detect_document_type(src_abs, content)
        stem = os.path.splitext(os.path.basename(src_abs))[0]
        bm_name = _sanitize_bookmark_name(f"chapter_{idx}_{stem}", f"chapter_{idx}")
        chapter_anchors[src_abs] = bm_name
        chapter_anchors[os.path.basename(src_abs)] = bm_name
        chapter_anchors[f"{stem}.md"] = bm_name
        chapter_anchors[f"{stem}.html"] = bm_name

        chap_title = _extract_title(content, os.path.basename(src_abs), True)
        if inferred_title is None:
            inferred_title = chap_title
        chapters_meta.append((src_abs, disp_name, content, bm_name, doc_info))

    toc_entries = [
        (bm_name, _extract_title(content, os.path.basename(src_abs), True))
        for src_abs, _, content, bm_name, _ in chapters_meta[1:]
    ]

    theme = parse_docx_theme(css_path)
    doc = docx.Document()
    configure_docx_document(doc, theme)

    meta_dt = resolve_office_datetime(timestamp=timestamp)
    doc.core_properties.title = inferred_title or "Drawlib Document"
    doc.core_properties.created = meta_dt
    doc.core_properties.modified = meta_dt

    cache = BuildImageCache(enabled=not no_cache)
    processor: DrawlibBlockProcessor | None = None
    total_files = len(chapters_meta)

    with tempfile.TemporaryDirectory() as scratch_dir:
        renderer = _DocxChapterRenderer(
            doc=doc,
            theme=theme,
            scratch_dir=scratch_dir,
            chapter_anchors=chapter_anchors,
        )

        for idx, (src_abs, disp_name, content, bm_name, doc_info) in enumerate(chapters_meta, start=1):
            if page_break and idx > 1:
                doc.add_page_break()

            total_steps = doc_info.block_count
            progress = FileBuildProgress(
                idx,
                total_files,
                file_name=disp_name,
                name_width=name_width,
                image_width=image_width,
            )
            progress.update(0, total_steps, done=False)

            if doc_info.has_drawlib and processor is None:
                processor = DrawlibBlockProcessor(
                    styles_path=styles_path,
                    utils_path=utils_path,
                    no_cache=no_cache,
                    cache=cache,
                    project_root=base_root,
                )

            src_dir = os.path.dirname(src_abs)
            orig_cwd = os.getcwd()
            sys_path_added = False
            try:
                os.chdir(src_dir)
                if src_dir not in sys.path:
                    sys.path.insert(0, src_dir)
                    sys_path_added = True

                renderer.render_chapter(
                    content,
                    src_abs=src_abs,
                    chapter_bookmark=bm_name,
                    processor=processor,
                    dpi=dpi,
                    progress_callback=progress.as_callback(),
                )
            finally:
                os.chdir(orig_cwd)
                if sys_path_added and src_dir in sys.path:
                    sys.path.remove(src_dir)

            progress.update(total_steps, total_steps, done=True)

            if idx == 1 and generate_index and toc_entries:
                _emit_toc_section(doc, theme, toc_entries, page_break=page_break)

        abs_output = os.path.abspath(output_path)
        out_dir = os.path.dirname(abs_output)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        doc.save(abs_output)

    if not timestamp:
        normalize_opc_zip_timestamps(abs_output)

    return file_list
