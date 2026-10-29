# Copyright (c) 2026 Yuichi Ito (yuichi@yuichi.com)
#
# This software is licensed under the Apache License, Version 2.0.
# For more information, please visit: https://github.com/yuichi110/drawlib
#
# This software is provided "as is", without warranty of any kind,
# express or implied, including but not limited to the warranties of
# merchantability, fitness for a particular purpose and noninfringement.

"""PowerPoint (.pptx) OpenXML writer from measured 1920x1080 slide DOM specs."""

from __future__ import annotations

import os
from typing import TYPE_CHECKING, Any, Final, Literal

from drawlib._builder._common.office import (
    ensure_static_office_image,
    normalize_opc_zip_timestamps,
    resolve_office_datetime,
)
from drawlib._slide._pptx_dom import (
    PptxCard,
    PptxColor,
    PptxDeckSpec,
    PptxImage,
    PptxParagraph,
    PptxRule,
    PptxRun,
    PptxSlideSpec,
    PptxTable,
    PptxTableCell,
    PptxTextBox,
    SlideElement,
)

if TYPE_CHECKING:
    from pptx.slide import Slide
    from pptx.table import _Cell
    from pptx.text.text import TextFrame, _Paragraph

_PX_TO_PT: Final[float] = 0.5
_STAGE_WIDTH_PX: Final[float] = 1920.0
_STAGE_WIDTH_PT: Final[float] = 960.0
_STAGE_HEIGHT_PT: Final[float] = 540.0

_RPR_EA_SUCCESSORS: Final[tuple[str, ...]] = (
    "a:cs",
    "a:sym",
    "a:hlinkClick",
    "a:hlinkMouseOver",
    "a:rtl",
    "a:extLst",
)
_RPR_HIGHLIGHT_SUCCESSORS: Final[tuple[str, ...]] = (
    "a:uLnTx",
    "a:uLn",
    "a:uFillTx",
    "a:uFill",
    "a:latin",
    "a:ea",
    *_RPR_EA_SUCCESSORS,
)
_TCPR_BORDER_SUCCESSORS: Final[tuple[str, ...]] = (
    "a:lnTlToBr",
    "a:lnBlToTr",
    "a:cell3D",
    "a:noFill",
    "a:solidFill",
    "a:gradFill",
    "a:blipFill",
    "a:pattFill",
    "a:grpFill",
    "a:headers",
    "a:extLst",
)


def _import_pptx() -> dict[str, Any]:
    """Lazily import python-pptx classes and enums."""
    try:
        from pptx import Presentation  # noqa: PLC0415
        from pptx.dml.color import RGBColor  # noqa: PLC0415
        from pptx.enum.shapes import MSO_SHAPE  # noqa: PLC0415
        from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_PARAGRAPH_ALIGNMENT  # noqa: PLC0415
        from pptx.oxml.ns import qn  # noqa: PLC0415
        from pptx.oxml.xmlchemy import OxmlElement  # noqa: PLC0415
        from pptx.util import Pt  # noqa: PLC0415
    except ImportError as err:
        msg = (
            "\n[PPTX Export Error]\n"
            "MS PowerPoint (.pptx) export requires the 'python-pptx' package.\n\n"
            "Reinstall drawlib or install python-pptx:\n"
            "  - For uv users:  uv sync  (or: uv add python-pptx)\n"
            "  - For pip users: pip install drawlib  (or: pip install python-pptx)\n"
        )
        raise RuntimeError(msg) from err

    return {
        "Presentation": Presentation,
        "RGBColor": RGBColor,
        "MSO_SHAPE": MSO_SHAPE,
        "MSO_VERTICAL_ANCHOR": MSO_VERTICAL_ANCHOR,
        "PP_PARAGRAPH_ALIGNMENT": PP_PARAGRAPH_ALIGNMENT,
        "qn": qn,
        "OxmlElement": OxmlElement,
        "Pt": Pt,
    }


def _hex6(rgb: tuple[int, int, int]) -> str:
    """Format an RGB tuple as a 6-digit uppercase hex string."""
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def _append_alpha_if_needed(srgb_el: Any, alpha: float, pmod: dict[str, Any]) -> None:  # noqa: ANN401
    """Append DrawingML <a:alpha> child when opacity is below 1.0."""
    if srgb_el is None or alpha >= 0.995:
        return
    qn = pmod["qn"]
    for existing in list(srgb_el.findall(qn("a:alpha"))):
        srgb_el.remove(existing)
    alpha_el = pmod["OxmlElement"]("a:alpha")
    alpha_val = max(0, min(100000, int(round(alpha * 100000))))
    alpha_el.set("val", str(alpha_val))
    srgb_el.append(alpha_el)


def _apply_fill_color(fill: Any, color: PptxColor | None, pmod: dict[str, Any]) -> None:  # noqa: ANN401
    """Apply solid RGB + optional alpha fill, or transparent background fill."""
    if color is None or color.alpha <= 0.01:
        fill.background()
        return
    fill.solid()
    fill.fore_color.rgb = pmod["RGBColor"](*color.rgb)
    if color.alpha < 0.995:
        qn = pmod["qn"]
        solid_fill = fill._xPr.find(qn("a:solidFill"))  # noqa: SLF001
        if solid_fill is not None:
            _append_alpha_if_needed(solid_fill.find(qn("a:srgbClr")), color.alpha, pmod)


def _apply_line_style(
    line: Any,  # noqa: ANN401
    color: PptxColor | None,
    width_px: float,
    pmod: dict[str, Any],
) -> None:
    """Apply outline border color, width, and optional alpha, or hide outline."""
    if color is None or width_px <= 0.0 or color.alpha <= 0.01:
        line.fill.background()
        return
    pt_fn = pmod["Pt"]
    line.color.rgb = pmod["RGBColor"](*color.rgb)
    line.width = pt_fn(max(0.5, width_px * _PX_TO_PT))
    if color.alpha < 0.995:
        qn = pmod["qn"]
        ln_el = line._get_or_add_ln()  # noqa: SLF001
        solid_fill = ln_el.find(qn("a:solidFill"))
        if solid_fill is not None:
            _append_alpha_if_needed(solid_fill.find(qn("a:srgbClr")), color.alpha, pmod)


def _set_slide_background(slide: Slide, bg: PptxColor, pmod: dict[str, Any]) -> None:
    """Apply solid slide background color."""
    background = slide.background
    _apply_fill_color(background.fill, bg, pmod)


def _draw_card(slide: Slide, card: PptxCard, pmod: dict[str, Any]) -> None:
    """Render a background/border container rectangle or rounded card."""
    pt_fn = pmod["Pt"]
    mso_shape = pmod["MSO_SHAPE"]
    shape_type = mso_shape.ROUNDED_RECTANGLE if card.radius_px > 0.5 else mso_shape.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type,
        pt_fn(card.x * _PX_TO_PT),
        pt_fn(card.y * _PX_TO_PT),
        pt_fn(max(1.0, card.width * _PX_TO_PT)),
        pt_fn(max(1.0, card.height * _PX_TO_PT)),
    )
    shape.shadow.inherit = False
    if card.radius_px > 0.5 and len(shape.adjustments) > 0:
        min_dim = max(1.0, min(card.width, card.height))
        shape.adjustments[0] = min(0.5, max(0.0, card.radius_px / min_dim))

    _apply_fill_color(shape.fill, card.bg, pmod)
    _apply_line_style(shape.line, card.border_color, card.border_width_px, pmod)

    if card.left_accent_color is not None and card.left_accent_width_px > 0:
        bar = slide.shapes.add_shape(
            mso_shape.RECTANGLE,
            pt_fn(card.x * _PX_TO_PT),
            pt_fn(card.y * _PX_TO_PT),
            pt_fn(max(1.5, card.left_accent_width_px * _PX_TO_PT)),
            pt_fn(max(1.0, card.height * _PX_TO_PT)),
        )
        bar.shadow.inherit = False
        _apply_fill_color(bar.fill, card.left_accent_color, pmod)
        bar.line.fill.background()


def _draw_rule(slide: Slide, rule: PptxRule, pmod: dict[str, Any]) -> None:
    """Render an <hr> divider line as a flat filled bar."""
    pt_fn = pmod["Pt"]
    mso_shape = pmod["MSO_SHAPE"]
    shape_type = mso_shape.ROUNDED_RECTANGLE if rule.radius_px > 0.5 else mso_shape.RECTANGLE
    bar = slide.shapes.add_shape(
        shape_type,
        pt_fn(rule.x * _PX_TO_PT),
        pt_fn(rule.y * _PX_TO_PT),
        pt_fn(max(1.0, rule.width * _PX_TO_PT)),
        pt_fn(max(1.0, rule.height * _PX_TO_PT)),
    )
    bar.shadow.inherit = False
    if rule.radius_px > 0.5 and len(bar.adjustments) > 0:
        bar.adjustments[0] = 0.5
    _apply_fill_color(bar.fill, rule.color, pmod)
    bar.line.fill.background()


def _draw_image(
    slide: Slide,
    img: PptxImage,
    scratch_dir: str,
    anim_frame: Literal["first", "last"],
    pmod: dict[str, Any],
) -> None:
    """Embed a high-DPI raster diagram or image at exact object-fit: contain coordinates."""
    pt_fn = pmod["Pt"]
    usable_path, _, _ = ensure_static_office_image(
        img.image_path,
        scratch_dir,
        anim_frame=anim_frame,
    )
    slide.shapes.add_picture(
        usable_path,
        pt_fn(img.x * _PX_TO_PT),
        pt_fn(img.y * _PX_TO_PT),
        width=pt_fn(max(1.0, img.width * _PX_TO_PT)),
        height=pt_fn(max(1.0, img.height * _PX_TO_PT)),
    )


def _compute_buffered_box_x_w(box: PptxTextBox) -> tuple[float, float]:
    """Compute alignment-preserving horizontal safety buffer to prevent substitute-font wrapping."""
    orig_x = box.x
    orig_w = max(4.0, box.width)
    extra = orig_w * (0.06 if box.is_code else 0.04) + (12.0 if box.is_code else 8.0)
    align = box.paragraphs[0].align if box.paragraphs else "left"

    if align == "center":
        max_sym = 2.0 * min(max(0.0, orig_x), max(0.0, _STAGE_WIDTH_PX - (orig_x + orig_w)))
        eff = min(extra, max_sym)
        return orig_x - eff / 2.0, orig_w + eff
    if align == "right":
        eff = min(extra, max(0.0, orig_x))
        return orig_x - eff, orig_w + eff
    return orig_x, min(orig_w + extra, max(orig_w, _STAGE_WIDTH_PX - orig_x))


def _apply_run_xml_extras(r_el: Any, run_spec: PptxRun, pmod: dict[str, Any]) -> None:  # noqa: ANN401
    """Attach strikethrough, inline code highlight, East Asian font, and color alpha on <a:rPr>."""
    oxml_element = pmod["OxmlElement"]
    qn = pmod["qn"]
    r_pr = r_el.get_or_add_rPr()

    if run_spec.strike:
        r_pr.set("strike", "sngStrike")

    if run_spec.color.alpha < 0.995:
        solid_fill = r_pr.find(qn("a:solidFill"))
        if solid_fill is not None:
            _append_alpha_if_needed(solid_fill.find(qn("a:srgbClr")), run_spec.color.alpha, pmod)

    if run_spec.code_bg is not None and run_spec.code_bg.alpha > 0.05:
        hl_el = oxml_element("a:highlight")
        srgb_el = oxml_element("a:srgbClr")
        srgb_el.set("val", _hex6(run_spec.code_bg.rgb))
        _append_alpha_if_needed(srgb_el, run_spec.code_bg.alpha, pmod)
        hl_el.append(srgb_el)
        r_pr.insert_element_before(hl_el, *_RPR_HIGHLIGHT_SUCCESSORS)  # noqa: SLF001

    if run_spec.ea_font_family:
        for old_ea in list(r_pr.findall(qn("a:ea"))):
            r_pr.remove(old_ea)
        ea_el = oxml_element("a:ea")
        ea_el.set("typeface", run_spec.ea_font_family)
        r_pr.insert_element_before(ea_el, *_RPR_EA_SUCCESSORS)  # noqa: SLF001


def _populate_run(p: _Paragraph, run_spec: PptxRun, pmod: dict[str, Any]) -> None:
    """Append a single styled inline run or line break to a PowerPoint paragraph."""
    if run_spec.is_line_break:
        p.add_line_break()
        return

    pt_fn = pmod["Pt"]
    r = p.add_run()
    r.text = run_spec.text
    if run_spec.href:
        r.hyperlink.address = run_spec.href
    r.font.name = run_spec.font_family
    r.font.size = pt_fn(max(4.0, run_spec.font_size_px * _PX_TO_PT))
    r.font.bold = run_spec.bold
    r.font.italic = run_spec.italic
    r.font.underline = run_spec.underline
    r.font.color.rgb = pmod["RGBColor"](*run_spec.color.rgb)
    _apply_run_xml_extras(r._r, run_spec, pmod)  # noqa: SLF001


def _configure_hanging_bullet(
    p: _Paragraph,
    para_spec: PptxParagraph,
    pmod: dict[str, Any],
) -> None:
    """Configure DrawingML hanging indent, tab stop, and leading bullet/number marker run."""
    if not para_spec.bullet_prefix or para_spec.hanging_indent_px <= 0:
        return

    pt_fn = pmod["Pt"]
    oxml_element = pmod["OxmlElement"]
    indent_len = pt_fn(para_spec.hanging_indent_px * _PX_TO_PT)
    indent_emu = int(indent_len)

    p_pr = p._p.get_or_add_pPr()  # noqa: SLF001
    p_pr.set("marL", str(indent_emu))
    p_pr.set("indent", str(-indent_emu))

    tab_lst = oxml_element("a:tabLst")
    tab = oxml_element("a:tab")
    tab.set("pos", str(indent_emu))
    tab_lst.append(tab)
    p_pr.insert_element_before(tab_lst, "a:defRPr", "a:extLst")  # noqa: SLF001

    ref_run = next((r for r in para_spec.runs if not r.is_line_break), None)
    marker_spec = PptxRun(
        text=f"{para_spec.bullet_prefix}\t",
        font_family=ref_run.font_family if ref_run else "Calibri",
        ea_font_family=ref_run.ea_font_family if ref_run else "Noto Sans JP",
        font_size_px=ref_run.font_size_px if ref_run else 20.0,
        color=ref_run.color if ref_run else PptxColor((32, 33, 36), 1.0),
        bold=False,
    )
    _populate_run(p, marker_spec, pmod)


def _populate_paragraph(
    p: _Paragraph,
    para_spec: PptxParagraph,
    pmod: dict[str, Any],
) -> None:
    """Configure paragraph alignment, line spacing, list marker, and inline runs."""
    pp_align = pmod["PP_PARAGRAPH_ALIGNMENT"]
    pt_fn = pmod["Pt"]
    align_map = {
        "left": pp_align.LEFT,
        "center": pp_align.CENTER,
        "right": pp_align.RIGHT,
        "justify": pp_align.JUSTIFY,
    }
    p.alignment = align_map.get(para_spec.align, pp_align.LEFT)

    # PowerPoint 1.0 line_spacing corresponds to ~1.2 CSS line-height
    ppt_line_spacing = max(0.95, min(1.85, round(para_spec.line_height_ratio / 1.2, 2)))
    p.line_spacing = ppt_line_spacing
    if para_spec.space_before_px > 0:
        p.space_before = pt_fn(para_spec.space_before_px * _PX_TO_PT)
    if para_spec.space_after_px > 0:
        p.space_after = pt_fn(para_spec.space_after_px * _PX_TO_PT)

    _configure_hanging_bullet(p, para_spec, pmod)
    for run_spec in para_spec.runs:
        _populate_run(p, run_spec, pmod)


def _populate_text_frame(
    tf: TextFrame,
    paragraphs: list[PptxParagraph],
    pmod: dict[str, Any],
) -> None:
    """Populate all paragraphs inside a PowerPoint TextFrame."""
    for idx, para_spec in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        _populate_paragraph(p, para_spec, pmod)


def _draw_textbox(slide: Slide, box: PptxTextBox, pmod: dict[str, Any]) -> None:
    """Render a native editable PowerPoint TextBox."""
    if not box.paragraphs:
        return
    pt_fn = pmod["Pt"]
    mso_anchor = pmod["MSO_VERTICAL_ANCHOR"]
    bx, bw = _compute_buffered_box_x_w(box)

    shape = slide.shapes.add_textbox(
        pt_fn(bx * _PX_TO_PT),
        pt_fn(box.y * _PX_TO_PT),
        pt_fn(max(4.0, bw * _PX_TO_PT)),
        pt_fn(max(4.0, box.height * _PX_TO_PT)),
    )
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = pt_fn(0)
    tf.margin_right = pt_fn(0)
    tf.margin_top = pt_fn(0)
    tf.margin_bottom = pt_fn(0)
    tf.vertical_anchor = mso_anchor.TOP
    _populate_text_frame(tf, box.paragraphs, pmod)


def _set_cell_border_side(
    tc_pr: Any,  # noqa: ANN401
    tag_name: str,
    border_spec: tuple[float, PptxColor] | None,
    pmod: dict[str, Any],
) -> None:
    """Set or clear a single table cell border side (a:lnL, a:lnR, a:lnT, a:lnB) on <a:tcPr>."""
    oxml_element = pmod["OxmlElement"]
    qn = pmod["qn"]
    pt_fn = pmod["Pt"]

    for existing in list(tc_pr.findall(qn(tag_name))):
        tc_pr.remove(existing)

    ln_el = oxml_element(tag_name)
    if border_spec is None:
        ln_el.append(oxml_element("a:noFill"))
    else:
        width_px, color = border_spec
        ln_el.set("w", str(int(pt_fn(max(0.5, width_px * _PX_TO_PT)))))
        solid_fill = oxml_element("a:solidFill")
        srgb = oxml_element("a:srgbClr")
        srgb.set("val", _hex6(color.rgb))
        _append_alpha_if_needed(srgb, color.alpha, pmod)
        solid_fill.append(srgb)
        ln_el.append(solid_fill)

    tc_pr.insert_element_before(ln_el, *_TCPR_BORDER_SUCCESSORS)  # noqa: SLF001


def _style_table_cell(cell: _Cell, spec: PptxTableCell, pmod: dict[str, Any]) -> None:
    """Apply padding, vertical alignment, fill, borders, and runs to a table cell."""
    pt_fn = pmod["Pt"]
    mso_anchor = pmod["MSO_VERTICAL_ANCHOR"]
    valign_map = {
        "top": mso_anchor.TOP,
        "middle": mso_anchor.MIDDLE,
        "bottom": mso_anchor.BOTTOM,
    }
    cell.margin_top = pt_fn(max(1.0, spec.pad_top_px * _PX_TO_PT))
    cell.margin_right = pt_fn(max(2.0, spec.pad_right_px * _PX_TO_PT))
    cell.margin_bottom = pt_fn(max(1.0, spec.pad_bottom_px * _PX_TO_PT))
    cell.margin_left = pt_fn(max(2.0, spec.pad_left_px * _PX_TO_PT))
    cell.vertical_anchor = valign_map.get(spec.valign, mso_anchor.MIDDLE)

    _apply_fill_color(cell.fill, spec.bg, pmod)

    tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
    _set_cell_border_side(tc_pr, "a:lnL", spec.border_left, pmod)
    _set_cell_border_side(tc_pr, "a:lnR", spec.border_right, pmod)
    _set_cell_border_side(tc_pr, "a:lnT", spec.border_top, pmod)
    _set_cell_border_side(tc_pr, "a:lnB", spec.border_bottom, pmod)

    tf = cell.text_frame
    tf.word_wrap = True
    _populate_text_frame(tf, spec.paragraphs, pmod)


def _draw_table(slide: Slide, tbl_spec: PptxTable, pmod: dict[str, Any]) -> None:
    """Render a native editable PowerPoint table."""
    n_rows = len(tbl_spec.rows)
    n_cols = len(tbl_spec.col_widths_px)
    if n_rows == 0 or n_cols == 0:
        return

    pt_fn = pmod["Pt"]
    shape = slide.shapes.add_table(
        n_rows,
        n_cols,
        pt_fn(tbl_spec.x * _PX_TO_PT),
        pt_fn(tbl_spec.y * _PX_TO_PT),
        pt_fn(max(10.0, tbl_spec.width * _PX_TO_PT)),
        pt_fn(max(10.0, tbl_spec.height * _PX_TO_PT)),
    )
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl.vert_banding = False
    tbl.last_row = False
    tbl.first_col = False
    tbl.last_col = False

    for c_idx, col_w in enumerate(tbl_spec.col_widths_px):
        tbl.columns[c_idx].width = pt_fn(max(8.0, col_w * _PX_TO_PT))
    for r_idx, row_h in enumerate(tbl_spec.row_heights_px):
        if r_idx < n_rows:
            tbl.rows[r_idx].height = pt_fn(max(8.0, row_h * _PX_TO_PT))

    for r_idx, row_cells in enumerate(tbl_spec.rows):
        for c_idx, cell_spec in enumerate(row_cells):
            if c_idx < n_cols:
                _style_table_cell(tbl.cell(r_idx, c_idx), cell_spec, pmod)


def _draw_slide_element(
    slide: Slide,
    elem: SlideElement,
    scratch_dir: str,
    anim_frame: Literal["first", "last"],
    pmod: dict[str, Any],
) -> None:
    """Dispatch a single SlideElement to its PowerPoint shape builder."""
    if isinstance(elem, PptxCard):
        _draw_card(slide, elem, pmod)
    elif isinstance(elem, PptxRule):
        _draw_rule(slide, elem, pmod)
    elif isinstance(elem, PptxImage):
        _draw_image(slide, elem, scratch_dir, anim_frame, pmod)
    elif isinstance(elem, PptxTable):
        _draw_table(slide, elem, pmod)
    elif isinstance(elem, PptxTextBox):
        _draw_textbox(slide, elem, pmod)


def _write_speaker_notes(
    slide: Slide,
    notes_text: str,
    default_ea_font: str,
    pmod: dict[str, Any],
) -> None:
    """Populate PowerPoint speaker notes with East Asian font support."""
    if not notes_text.strip():
        return
    notes_tf = slide.notes_slide.notes_text_frame
    if notes_tf is None:
        return
    notes_tf.text = notes_text

    oxml_element = pmod["OxmlElement"]
    qn = pmod["qn"]
    for para in notes_tf.paragraphs:
        for run in para.runs:
            r_pr = run._r.get_or_add_rPr()  # noqa: SLF001
            for old_ea in list(r_pr.findall(qn("a:ea"))):
                r_pr.remove(old_ea)
            ea_el = oxml_element("a:ea")
            ea_el.set("typeface", default_ea_font)
            r_pr.insert_element_before(ea_el, *_RPR_EA_SUCCESSORS)  # noqa: SLF001


def write_deck_to_pptx(
    deck: PptxDeckSpec,
    output_path: str,
    scratch_dir: str,
    *,
    title: str | None = None,
    anim_frame: Literal["first", "last"] = "last",
    timestamp: bool = False,
) -> str:
    """Write a measured PptxDeckSpec to a 16:9 widescreen .pptx file.

    Args:
        deck: Measured slide deck layout specification.
        output_path: Destination .pptx path.
        scratch_dir: Temporary directory for any converted static raster images.
        title: Optional presentation metadata title override.
        anim_frame: Frame selection ('first' or 'last') for multi-frame animations.
        timestamp: If True, keep current build timestamp; otherwise normalize deterministically.

    Returns:
        str: Absolute path to the saved .pptx file.
    """
    pmod = _import_pptx()
    pt_fn = pmod["Pt"]
    prs = pmod["Presentation"]()

    # Exact 16:9 Widescreen 960pt x 540pt (12,192,000 x 6,858,000 EMU)
    prs.slide_width = pt_fn(_STAGE_WIDTH_PT)
    prs.slide_height = pt_fn(_STAGE_HEIGHT_PT)
    blank_layout = prs.slide_layouts[6]

    for slide_spec in deck.slides:
        _render_single_slide(
            prs,
            blank_layout,
            slide_spec,
            deck.default_ea_font,
            scratch_dir,
            anim_frame,
            pmod,
        )

    dt = resolve_office_datetime(timestamp=timestamp)
    props = prs.core_properties
    props.title = title or deck.title or "Drawlib Presentation"
    props.author = "drawlib"
    props.comments = ""
    props.created = dt
    props.modified = dt

    abs_out = os.path.abspath(output_path)
    out_dir = os.path.dirname(abs_out)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    prs.save(abs_out)

    if not timestamp:
        normalize_opc_zip_timestamps(abs_out)

    return abs_out


def _render_single_slide(
    prs: Any,  # noqa: ANN401
    blank_layout: Any,  # noqa: ANN401
    slide_spec: PptxSlideSpec,
    default_ea_font: str,
    scratch_dir: str,
    anim_frame: Literal["first", "last"],
    pmod: dict[str, Any],
) -> None:
    """Create and populate one PowerPoint slide from a PptxSlideSpec."""
    slide = prs.slides.add_slide(blank_layout)
    _set_slide_background(slide, slide_spec.bg, pmod)
    for elem in slide_spec.elements:
        _draw_slide_element(slide, elem, scratch_dir, anim_frame, pmod)
    _write_speaker_notes(slide, slide_spec.notes_text, default_ea_font, pmod)
