# Copyright (c) 2026 Yuichi Ito (yuichi@yuichi.com)
#
# This software is licensed under the Apache License, Version 2.0.
# For more information, please visit: https://github.com/yuichi110/drawlib
#
# This software is provided "as is", without warranty of any kind,
# express or implied, including but not limited to the warranties of
# merchantability, fitness for a particular purpose and noninfringement.

"""Unit and integration tests for MS Word (.docx) document export."""

from __future__ import annotations

import datetime
import zipfile

import docx
import pytest
from docx.oxml.ns import qn

from drawlib._builder.doc_builder import build_document, build_docx


def test_build_docx_rich_markdown_elements(tmp_path) -> None:
    """Verify headings, inline styles, links, lists, blockquotes, tables, and code blocks."""
    doc_dir = tmp_path / "spec_src"
    doc_dir.mkdir()

    (doc_dir / "style.css").write_text(
        """:root {
    --rtd-brand-color: #1a73e8;
    --rtd-heading-color: #202124;
    --rtd-text-color: #3c4043;
    --rtd-bg-color: #ffffff;
    --rtd-body-font: "Roboto", "Noto Sans JP", sans-serif;
}
""",
        encoding="utf-8",
    )

    (doc_dir / "00_cover.md").write_text(
        """# System Architecture Specification

Lead paragraph with **bold**, *italic*, ~~strikethrough~~, `inline_code()`, and [Drawlib Site](https://www.drawlib.com).

See [Chapter 1](01_details.md) for implementation details.

> **Important**: All endpoints require mTLS authentication.

---

- Top item A
  - Nested bullet A1
- Top item B

1. First step
2. Second step

Paragraph between numbered lists.

1. Restarted first step
2. Restarted second step
""",
        encoding="utf-8",
    )

    (doc_dir / "01_details.md").write_text(
        """# Implementation Details

## Service Matrix

| Service | Port | Status |
| :--- | :---: | ---: |
| `gateway` | 8443 | **Active** |
| `worker` | 9090 | Standby |

```python
def health_check() -> bool:
    return True
```
""",
        encoding="utf-8",
    )

    out_docx = tmp_path / "spec.docx"
    result_path = build_docx(
        input_path=str(doc_dir),
        output_file=str(out_docx),
        generate_index=True,
        page_break=True,
        no_cache=True,
    )

    assert result_path == str(out_docx)
    assert out_docx.is_file()

    parsed = docx.Document(str(out_docx))
    assert parsed.core_properties.title == "System Architecture Specification"
    assert parsed.core_properties.created == datetime.datetime(2026, 1, 1, 0, 0, 0, tzinfo=datetime.timezone.utc)

    # Verify A4 geometry (210mm x 297mm)
    sec = parsed.sections[0]
    assert sec.page_width is not None
    assert sec.page_height is not None
    assert abs(sec.page_width.mm - 210.0) < 0.2
    assert abs(sec.page_height.mm - 297.0) < 0.2

    # Verify East Asian font w:eastAsia is set on Normal style
    normal_rpr = parsed.styles["Normal"].element.rPr
    assert normal_rpr is not None
    normal_rfonts = normal_rpr.find(qn("w:rFonts"))
    assert normal_rfonts is not None
    assert normal_rfonts.get(qn("w:eastAsia")) == "Noto Sans JP"

    # Verify Heading paragraphs and Table of Contents
    heading_texts = [p.text for p in parsed.paragraphs if p.style and p.style.name.startswith("Heading")]
    assert "System Architecture Specification" in heading_texts
    assert "Table of Contents" in heading_texts
    assert "Implementation Details" in heading_texts
    assert "Service Matrix" in heading_texts

    # Verify hyperlinks (external relationship + internal w:anchor bookmarks)
    xml_body = parsed.element.body.xml
    assert "w:hyperlink" in xml_body
    assert 'w:anchor="chapter_2_01_details"' in xml_body
    assert "w:shd" in xml_body  # inline code / callout shading
    assert "w:bookmarkStart" in xml_body

    # Verify GFM table + 1x1 syntax-highlighted code block table exist
    assert len(parsed.tables) == 2
    gfm_tbl = parsed.tables[0]
    assert len(gfm_tbl.rows) == 3
    assert len(gfm_tbl.columns) == 3
    assert gfm_tbl.cell(0, 0).text.strip() == "Service"
    assert gfm_tbl.cell(1, 0).text.strip() == "gateway"
    tr_pr = gfm_tbl.rows[0]._tr.trPr
    assert tr_pr is not None
    assert tr_pr.find(qn("w:tblHeader")) is not None

    code_tbl = parsed.tables[1]
    assert "def health_check" in code_tbl.cell(0, 0).text


def test_build_docx_drawlib_blocks_and_animation(tmp_path) -> None:
    """Verify code=fold/hide omit code while code=show includes code, and animations flatten to PNG."""
    md_file = tmp_path / "diagrams.md"
    md_file.write_text(
        """# Diagram Export Modes

```drawlib code=fold file=folded_fig.png caption="Figure 1: Folded Diagram"
from drawlib.canvas import setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

setup(width=80, height=30)
rectangle((40, 15), width=50, height=16, style=Styles.Neutral, text="Folded")
```

```drawlib code=show width=80% align=left file=shown_fig.png caption="Figure 2: Shown Code Diagram"
from drawlib.canvas import setup
from drawlib.shapes import circle
from drawlib.styles import Styles

setup(width=60, height=30)
circle((30, 15), radius=10, style=Styles.PrimaryFlat)
```

```drawlib file=anim_fig.png caption="Figure 3: Animated Final Frame"
from drawlib.anim import Animation
from drawlib.canvas import save, setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

setup(width=60, height=30)
anim = Animation(fps=2)
with anim.frame():
    rectangle((20, 15), width=15, height=12, style=Styles.Neutral)
with anim.frame():
    rectangle((40, 15), width=25, height=16, style=Styles.SecondaryFlat)
save()
```
""",
        encoding="utf-8",
    )

    out_docx = build_docx(str(md_file), dpi=160, no_cache=True)
    parsed = docx.Document(out_docx)

    # Exactly 1 code table should exist (only for code=show, NOT code=fold or default code=hide)
    assert len(parsed.tables) == 1
    assert "circle((30, 15)" in parsed.tables[0].cell(0, 0).text
    assert "Folded" not in parsed.tables[0].cell(0, 0).text

    # All 3 diagrams should be embedded as inline shapes and captions
    assert len(parsed.inline_shapes) == 3
    all_para_text = "\n".join(p.text for p in parsed.paragraphs)
    assert "Figure 1: Folded Diagram" in all_para_text
    assert "Figure 2: Shown Code Diagram" in all_para_text
    assert "Figure 3: Animated Final Frame" in all_para_text


def test_build_docx_deterministic_and_build_document_dispatch(tmp_path) -> None:
    """Verify deterministic byte-identical output and build_document dispatch."""
    src_dir = tmp_path / "doc_src"
    src_dir.mkdir()
    (src_dir / "00_intro.md").write_text("# Deterministic Word\n\n日本語テキストの検証。", encoding="utf-8")

    out1 = tmp_path / "out1.docx"
    out2 = tmp_path / "out2.docx"
    build_docx(str(src_dir), str(out1))
    build_document(str(src_dir), str(out2))  # inferred from .docx suffix

    assert out1.read_bytes() == out2.read_bytes()
    with zipfile.ZipFile(out1, "r") as zf:
        for info in zf.infolist():
            assert info.date_time == (2026, 1, 1, 0, 0, 0)


def test_build_docx_dark_theme_fallback_and_slide_rejection(tmp_path) -> None:
    """Verify dark CSS theme falls back to readable text colors and slide directories are rejected."""
    dark_dir = tmp_path / "dark_doc_src"
    dark_dir.mkdir()
    (dark_dir / "style.css").write_text(
        ":root { --rtd-bg-color: #0f172a; --rtd-text-color: #f8fafc; --rtd-brand-color: #38bdf8; }",
        encoding="utf-8",
    )
    (dark_dir / "01_page.md").write_text("# Dark Theme Chapter\n\nBody text.", encoding="utf-8")
    dark_out = build_docx(str(dark_dir), str(tmp_path / "dark.docx"))
    parsed = docx.Document(dark_out)
    # Normal text color must stay dark on white A4 Word paper, not #f8fafc near-white
    rgb = parsed.styles["Normal"].font.color.rgb
    assert rgb is not None
    assert rgb[0] < 100 and rgb[1] < 100 and rgb[2] < 100

    # Slide directory must raise clear error suggesting drawlib build pptx
    slide_dir = tmp_path / "slide_src"
    slide_dir.mkdir()
    (slide_dir / "slide.css").write_text(".presentation-stage {}", encoding="utf-8")
    (slide_dir / "01_title.md").write_text("# Slide", encoding="utf-8")

    with pytest.raises(ValueError, match="drawlib build pptx"):
        build_docx(str(slide_dir))
