# Copyright (c) 2026 Yuichi Ito (yuichi@yuichi.com)
#
# This software is licensed under the Apache License, Version 2.0.
# For more information, please visit: https://github.com/yuichi110/drawlib
#
# This software is provided "as is", without warranty of any kind,
# express or implied, including but not limited to the warranties of
# merchantability, fitness for a particular purpose and noninfringement.

"""Unit and integration tests for MS PowerPoint (.pptx) slide deck export."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from drawlib.slide import build_pptx


def _collect_all_slide_text(slide) -> str:
    """Concatenate all text runs from text frames and tables on a slide."""
    chunks: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            chunks.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    chunks.append(cell.text)
    return "\n".join(chunks)


class TestBuildPptx:
    """Test suite for slide deck -> MS PowerPoint (.pptx) compilation."""

    def test_widescreen_geometry_typography_tables_code_and_notes(self, tmp_path: Path) -> None:
        """Verify 0.5x coordinate mapping, rich text, lists, tables, code blocks, CJK fonts, and speaker notes."""
        slide_dir = tmp_path / "deck_src"
        slide_dir.mkdir()

        (slide_dir / "01_overview.md").write_text(
            """\
::: block xy=(80, 40) width=1760 height=90
# アーキテクチャ概要 (Architecture Overview)
:::

::: block xy=(80, 160) width=820 height=780
Intro with **bold**, *italic*, ~~strike~~, `inline_code`, and [Docs](https://example.com).

- First bullet item
- Second bullet item

1. Step one
2. Step two

---

```python
def greet(name: str) -> str:
    return f"Hello {name}"
```
:::

::: block xy=(940, 160) width=900 height=420
| Service | Latency | Status |
| :--- | :---: | ---: |
| Gateway | 4.2ms | Healthy |
| Primary DB | 1.8ms | Replicated |
:::

::: block xy=(940, 600) width=900 height=340
```drawlib file=overview_diagram.svg
from drawlib.canvas import clear, save, setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

clear()
setup(width=90, height=34)
rectangle((45, 17), width=70, height=22, style=Styles.PrimaryFlat, text="Core Service")
save()
```
:::

::: note
発表者向けのスピーカーノートです。

- ポイント1: レイテンシ目標の達成
- ポイント2: 冗長構成の確認
:::
""",
            encoding="utf-8",
        )

        out_pptx = build_pptx(str(slide_dir), dpi=200)
        assert Path(out_pptx) == tmp_path / "deck.pptx"
        assert Path(out_pptx).is_file()

        prs = Presentation(out_pptx)
        assert prs.slide_width == Pt(960)
        assert prs.slide_height == Pt(540)
        assert len(prs.slides) == 1

        slide = prs.slides[0]
        full_text = _collect_all_slide_text(slide)
        assert "アーキテクチャ概要" in full_text
        assert "inline_code" in full_text
        assert "•\t" in full_text
        assert "1.\t" in full_text
        assert "def" in full_text
        assert "greet" in full_text

        # Verify heading coordinate mapping (~0.5x of xy=(80, 40) -> left=40pt, top=20pt)
        heading_shapes = [
            s for s in slide.shapes if s.has_text_frame and "アーキテクチャ概要" in s.text_frame.text
        ]
        assert len(heading_shapes) == 1
        assert abs(heading_shapes[0].left.pt - 40.0) < 4.0
        assert abs(heading_shapes[0].top.pt - 20.0) < 6.0

        # Verify native PowerPoint table
        tables = [s.table for s in slide.shapes if s.has_table]
        assert len(tables) == 1
        tbl = tables[0]
        assert len(tbl.rows) == 3
        assert len(tbl.columns) == 3
        assert tbl.cell(0, 0).text.strip() == "Service"
        assert tbl.cell(1, 0).text.strip() == "Gateway"
        assert tbl.cell(2, 2).text.strip() == "Replicated"

        # Verify hyperlink + strict DrawingML <a:rPr> ordering (<a:ea> before <a:hlinkClick>)
        hlink_runs = []
        ea_found = False
        strike_found = False
        highlight_found = False
        for s in slide.shapes:
            if not s.has_text_frame:
                continue
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    r_pr = r._r.rPr
                    if r_pr is not None:
                        ea_el = r_pr.find(qn("a:ea"))
                        if ea_el is not None and ea_el.get("typeface"):
                            ea_found = True
                        if r_pr.get("strike") == "sngStrike":
                            strike_found = True
                        if r_pr.find(qn("a:highlight")) is not None:
                            highlight_found = True
                        if r.hyperlink and r.hyperlink.address:
                            hlink_runs.append(r.hyperlink.address)
                            tags = [child.tag for child in r_pr]
                            assert tags.index(qn("a:ea")) < tags.index(qn("a:hlinkClick"))

        assert "https://example.com" in hlink_runs
        assert ea_found
        assert strike_found
        assert highlight_found

        # Verify embedded high-DPI PNG diagram (explicit .svg rewritten to high-DPI .png)
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == 1
        with Image.open(io.BytesIO(pictures[0].image.blob)) as img:
            assert img.format == "PNG"
            assert img.width >= 1500  # 90 units -> 9 inches * 200 DPI = 1800 px

        # Verify speaker notes
        notes_tf = slide.notes_slide.notes_text_frame
        assert notes_tf is not None
        assert "発表者向けのスピーカーノートです。" in notes_tf.text
        assert "• ポイント1: レイテンシ目標の達成" in notes_tf.text

    def test_z_index_layering_and_rgba_floating_card(self, tmp_path: Path) -> None:
        """Verify z_index stacking order across blocks and DrawingML alpha on rgba() cards."""
        slide_dir = tmp_path / "zdeck_src"
        slide_dir.mkdir()

        card_css = (
            "background: rgba(15, 23, 42, 0.85); color: #f8fafc; "
            "padding: 20px; border-radius: 16px; border: 2px solid #6366f1;"
        )
        (slide_dir / "01_layers.md").write_text(
            f"""\
::: block xy=(1680, 990) width=160 height=40 z_index=20
```drawlib file=badge_top.svg
from drawlib.canvas import clear, save, setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

clear()
setup(width=16, height=4)
rectangle((8, 2), width=15, height=3.5, style=Styles.AccentFlat, text="TOP")
save()
```
:::

::: block xy=(900, 0) width=1020 height=1080 z_index=2
```drawlib file=bg_canvas.svg
from drawlib.canvas import clear, save, setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

clear()
setup(width=102, height=108)
rectangle((51, 54), width=102, height=108, style=Styles.Neutral)
save()
```
:::

::: block xy=(740, 720) width=520 height=160 z_index=15 style="{card_css}"
**Floating Glass Card**
:::
""",
            encoding="utf-8",
        )

        out_pptx = tmp_path / "layers.pptx"
        build_pptx(str(slide_dir), str(out_pptx), dpi=120)

        prs = Presentation(str(out_pptx))
        slide = prs.slides[0]
        shapes = list(slide.shapes)

        # Find indices of background picture (~510pt), floating card (~260pt), text, and top badge (~80pt)
        bg_pic_idx = next(
            i for i, s in enumerate(shapes) if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width.pt > 400
        )
        top_pic_idx = next(
            i for i, s in enumerate(shapes) if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width.pt < 150
        )
        card_idx = next(
            i
            for i, s in enumerate(shapes)
            if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and abs(s.left.pt - 370.0) < 5.0
        )
        text_idx = next(
            i for i, s in enumerate(shapes) if s.has_text_frame and "Floating Glass Card" in s.text_frame.text
        )

        assert bg_pic_idx < card_idx < text_idx < top_pic_idx

        # Check DrawingML <a:alpha val="85000"/> on the rgba(15, 23, 42, 0.85) card
        card_shape = shapes[card_idx]
        alpha_els = card_shape._element.findall(".//" + qn("a:alpha"))
        assert any(el.get("val") == "85000" for el in alpha_els)

    def test_animation_first_vs_last_frame_extraction(self, tmp_path: Path) -> None:
        """Verify Animation() blocks flatten cleanly to the requested 'first' or 'last' static frame."""
        slide_dir = tmp_path / "anim_src"
        slide_dir.mkdir()

        (slide_dir / "01_anim.md").write_text(
            """\
::: block xy=(160, 140) width=1600 height=800
```drawlib file=motion.png anim_trigger=auto
from drawlib.anim import Animation
from drawlib.canvas import clear, save, setup
from drawlib.shapes import rectangle
from drawlib.styles import Styles

clear()
setup(width=80, height=40)
anim = Animation(fps=4.0, loop=0)

with anim.frame(duration=0.2):
    rectangle((40, 20), width=80, height=40, style=Styles.WhiteFlat.patch(shape_fill_color=(220, 30, 30)))

with anim.frame(duration=0.2):
    rectangle((40, 20), width=80, height=40, style=Styles.WhiteFlat.patch(shape_fill_color=(30, 80, 220)))

save()
```
:::
""",
            encoding="utf-8",
        )

        first_pptx = tmp_path / "first.pptx"
        last_pptx = tmp_path / "last.pptx"
        build_pptx(str(slide_dir), str(first_pptx), dpi=100, anim_frame="first")
        build_pptx(str(slide_dir), str(last_pptx), dpi=100, anim_frame="last")

        def _center_rgb(pptx_file: Path) -> tuple[int, int, int]:
            prs = Presentation(str(pptx_file))
            pics = [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            assert len(pics) == 1
            with Image.open(io.BytesIO(pics[0].image.blob)) as im:
                assert not getattr(im, "is_animated", False)
                rgba = im.convert("RGB")
                return rgba.getpixel((rgba.width // 2, rgba.height // 2))

        r_first, _, b_first = _center_rgb(first_pptx)
        r_last, _, b_last = _center_rgb(last_pptx)
        assert r_first > 180 and b_first < 80
        assert b_last > 180 and r_last < 80

    def test_deterministic_pptx_archive_and_validation(
        self,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Verify deterministic OPC archive output and input validation errors."""
        monkeypatch.setenv("SOURCE_DATE_EPOCH", "1735689600")
        slide_dir = tmp_path / "det_src"
        slide_dir.mkdir()
        (slide_dir / "01_title.md").write_text("# Deterministic Deck\n\nReproducible output.", encoding="utf-8")

        out_a = tmp_path / "a.pptx"
        out_b = tmp_path / "b.pptx"
        build_pptx(str(slide_dir), str(out_a), timestamp=False)
        build_pptx(str(slide_dir), str(out_b), timestamp=False)
        assert out_a.read_bytes() == out_b.read_bytes()

        with pytest.raises(ValueError, match="Input directory does not exist"):
            build_pptx(str(tmp_path / "missing_dir"))

        with pytest.raises(ValueError, match="dpi must be a positive integer"):
            build_pptx(str(slide_dir), dpi=0)
